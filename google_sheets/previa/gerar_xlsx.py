"""
Monta Robo_LPCO_previa.xlsx a partir de saida/modelo.json (gerado por gerar_previa.mjs).

O arquivo pode ser importado no Google Drive para ver as abas como ficam depois de uma
atualização e de uma busca, com dados ficticios. Ele mostra o resultado, mas nao tem o
script.

Uso: python google_sheets/previa/gerar_xlsx.py
"""

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

PASTA = Path(__file__).resolve().parent
ENTRADA = PASTA / "saida" / "modelo.json"
SAIDA = PASTA / "Robo_LPCO_previa.xlsx"
PIXELS_POR_LARGURA = 7


def cor(hexa):
    return hexa.lstrip("#").upper() if hexa else None


def montar(abas):
    wb = Workbook()
    wb.remove(wb.active)
    for aba in abas:
        ws = wb.create_sheet(aba["nome"][:31])
        if aba["cor"]:
            ws.sheet_properties.tabColor = cor(aba["cor"])
        listas = {}
        for c in aba["celulas"]:
            cel = ws.cell(row=c["l"], column=c["c"], value=c["v"] if c["v"] != "" else None)
            if c["fundo"]:
                cel.fill = PatternFill("solid", fgColor=cor(c["fundo"]))
            cel.font = Font(name="Arial", bold=c["negrito"], italic=c["italico"],
                            color=cor(c["cor"]) or "202124")
            cel.alignment = Alignment(vertical="top", wrap_text=c["quebra"])
            if c["lista"]:
                listas.setdefault(tuple(c["lista"]), []).append(cel.coordinate)
        for opcoes, celulas in listas.items():
            dv = DataValidation(type="list", formula1='"' + ",".join(opcoes) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            for coord in celulas:
                dv.add(coord)
        for col, largura in aba["larguras"].items():
            ws.column_dimensions[get_column_letter(int(col))].width = max(int(largura) / PIXELS_POR_LARGURA, 6)
        for col in aba["ocultas"]:
            ws.column_dimensions[get_column_letter(int(col))].hidden = True
        if aba["oculta"]:
            ws.sheet_state = "hidden"
    wb.active = wb.sheetnames.index("Buscar")
    wb.save(SAIDA)
    return SAIDA


if __name__ == "__main__":
    saida = montar(json.loads(ENTRADA.read_text(encoding="utf-8")))
    print(f"{saida.name} gravado")
