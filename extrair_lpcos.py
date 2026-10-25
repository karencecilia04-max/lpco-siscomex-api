"""
Extrator de LPCOs via API do Portal Unico Siscomex.

Lista as LPCOs de importacao registradas no CNPJ (GET /talpco/api/ext/lpco/consulta),
detalha cada uma (GET /talpco/api/ext/lpco/{numero}) e gera um Excel com:
numero da LPCO, item, NCM e codigo/versao do produto do Catalogo de Produtos.

Configuracao: config_lpco.json (mesma pasta). Veja config_lpco.exemplo.json.
Autenticacao: chave de acesso (Client-Id/Client-Secret gerados no portal)
ou certificado digital A1 (.pfx). Certificado A3 (token/cartao) nao e suportado.
"""

import getpass
import json
import os
import re
import sys
import tempfile
import time
from datetime import date, datetime, timedelta

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PASTA = os.path.dirname(os.path.abspath(__file__))
ARQ_CONFIG = os.path.join(PASTA, "config_lpco.json")
ARQ_CACHE = os.path.join(PASTA, "cache_detalhes_lpco.json")

AMBIENTES = {
    "producao": "https://portalunico.siscomex.gov.br",
    "validacao": "https://val.portalunico.siscomex.gov.br",
}


class LimiteAtingido(Exception):
    """Portal bloqueou por excesso de acessos na hora (PUCX-ER1001)."""


def log(msg):
    print(f"[{datetime.now():%H:%M:%S}] {msg}", flush=True)


def carregar_config():
    if not os.path.exists(ARQ_CONFIG):
        sys.exit(f"Arquivo {ARQ_CONFIG} nao encontrado. Copie config_lpco.exemplo.json "
                 "para config_lpco.json e preencha.")
    with open(ARQ_CONFIG, encoding="utf-8") as f:
        cfg = json.load(f)
    cfg["cnpj"] = "".join(c for c in str(cfg["cnpj"]) if c.isdigit())
    if len(cfg["cnpj"]) != 14:
        sys.exit("CNPJ no config deve ter 14 digitos.")
    if cfg.get("ambiente", "producao") not in AMBIENTES:
        sys.exit(f"ambiente deve ser um de: {', '.join(AMBIENTES)}")
    return cfg


# ---------------------------------------------------------------------------
# Autenticacao
# ---------------------------------------------------------------------------

def pfx_para_pem(caminho_pfx, senha):
    """Converte o .pfx em arquivos PEM temporarios (requests nao le .pfx direto)."""
    from cryptography.hazmat.primitives.serialization import (
        Encoding, NoEncryption, PrivateFormat, pkcs12)

    with open(caminho_pfx, "rb") as f:
        chave, cert, extras = pkcs12.load_key_and_certificates(f.read(), senha.encode())
    pasta_tmp = tempfile.mkdtemp(prefix="lpco_cert_")
    arq_cert = os.path.join(pasta_tmp, "cert.pem")
    arq_chave = os.path.join(pasta_tmp, "chave.pem")
    with open(arq_cert, "wb") as f:
        f.write(cert.public_bytes(Encoding.PEM))
        for c in extras or []:
            f.write(c.public_bytes(Encoding.PEM))
    with open(arq_chave, "wb") as f:
        f.write(chave.private_bytes(Encoding.PEM, PrivateFormat.PKCS8, NoEncryption()))
    return arq_cert, arq_chave, pasta_tmp


class PortalUnico:
    def __init__(self, cfg):
        self.base = AMBIENTES[cfg.get("ambiente", "producao")]
        self.cfg = cfg
        self.pausa = float(cfg.get("pausa_entre_chamadas_seg", 1.5))
        self.sessao = requests.Session()
        self.pasta_cert = None
        auth = cfg["autenticacao"]
        if auth["tipo"] == "certificado":
            senha = auth.get("senha_pfx") or getpass.getpass("Senha do certificado .pfx: ")
            arq_cert, arq_chave, self.pasta_cert = pfx_para_pem(auth["arquivo_pfx"], senha)
            self.sessao.cert = (arq_cert, arq_chave)
        elif auth["tipo"] != "chave_acesso":
            sys.exit("autenticacao.tipo deve ser 'chave_acesso' ou 'certificado'.")
        self.autenticar()

    def autenticar(self):
        auth = self.cfg["autenticacao"]
        headers = {"Role-Type": self.cfg.get("perfil", "IMPEXP")}
        if auth["tipo"] == "chave_acesso":
            url = f"{self.base}/portal/api/autenticar/chave-acesso"
            headers["Client-Id"] = auth["client_id"]
            headers["Client-Secret"] = auth["client_secret"]
        else:
            url = f"{self.base}/portal/api/autenticar"
        r = self.sessao.post(url, headers=headers, timeout=60)
        if r.status_code != 200:
            sys.exit(f"Falha na autenticacao ({r.status_code}): {r.text[:500]}")
        self.sessao.headers["Authorization"] = r.headers["Set-Token"]
        self.sessao.headers["X-CSRF-Token"] = r.headers["X-CSRF-Token"]
        log(f"Autenticado em {self.base} (perfil {headers['Role-Type']}).")

    def _get_com_retentativa(self, url, params):
        """Repete a chamada se a conexao cair ou o portal der erro 5xx momentaneo."""
        esperas = [10, 30, 60, 120]
        for n, espera in enumerate(esperas + [None], 1):
            time.sleep(self.pausa)
            try:
                r = self.sessao.get(url, params=params, timeout=120)
                if r.status_code < 500 or espera is None:
                    return r
                motivo = f"erro {r.status_code} do portal"
            except (requests.ConnectionError, requests.Timeout) as e:
                if espera is None:
                    raise
                motivo = f"conexao caiu ({type(e).__name__})"
            log(f"{motivo}; tentativa {n}/{len(esperas)}, aguardando {espera}s...")
            time.sleep(espera)

    def get(self, caminho, params=None, sistema="talpco", binario=False):
        url = f"{self.base}/{sistema}/api{caminho}"
        for tentativa in range(2):
            r = self._get_com_retentativa(url, params)
            if r.status_code == 401 and tentativa == 0:
                log("Token expirado, autenticando de novo...")
                self.autenticar()
                continue
            break
        if "PUCX-ER1001" in r.text:
            raise LimiteAtingido(r.json().get("message", r.text))
        if r.status_code == 204:
            return r.status_code, None
        if r.status_code not in (200, 206):
            raise RuntimeError(f"GET {caminho} -> {r.status_code}: {r.text[:500]}")
        return r.status_code, (r.content if binario else r.json())

    def fechar(self):
        if self.pasta_cert:
            for nome in os.listdir(self.pasta_cert):
                os.remove(os.path.join(self.pasta_cert, nome))
            os.rmdir(self.pasta_cert)


# ---------------------------------------------------------------------------
# Consulta
# ---------------------------------------------------------------------------

def janelas_de_datas(inicio, fim, dias):
    """Quebra o periodo em janelas menores para nao estourar a consulta."""
    atual = inicio
    while atual <= fim:
        final = min(atual + timedelta(days=dias - 1), fim)
        yield atual, final
        atual = final + timedelta(days=1)


def ncms_do_config(cfg):
    """NCMs do filtro, so digitos (aceita "4011.10.00" ou "40111000"). Lista vazia = todas."""
    ncms = ["".join(c for c in str(n) if c.isdigit()) for n in cfg.get("ncms") or []]
    for n in ncms:
        if len(n) != 8:
            sys.exit(f"NCM invalida no config: {n} (precisa ter 8 digitos).")
    return ncms


def listar_lpcos(api, cfg):
    inicio = date.fromisoformat(cfg["data_inicial"])
    fim = date.fromisoformat(cfg.get("data_final") or date.today().isoformat())
    ncms = ncms_do_config(cfg)
    encontrados = {}
    for ini, fin in janelas_de_datas(inicio, fim, int(cfg.get("dias_por_janela", 30))):
        # Uma consulta por NCM (ncm-inicial = ncm-final), para nao pegar NCMs no meio da faixa
        for ncm in ncms or [None]:
            offset = 0
            while True:
                params = {
                    "tipo-operacao": "IMPORTACAO",
                    "importador-exportador": cfg["cnpj"],
                    "data-inicial-registro": ini.isoformat(),
                    "data-final-registro": fin.isoformat(),
                    "offset": offset,
                }
                if cfg.get("situacao"):
                    params["situacao"] = cfg["situacao"]
                if ncm:
                    params["ncm-inicial"] = ncm
                    params["ncm-final"] = ncm
                _, dados = api.get("/ext/lpco/consulta", params)
                if not dados:
                    break
                novos = [d for d in dados if d["numero"] not in encontrados]
                for d in novos:
                    encontrados[d["numero"]] = d
                # Para se a pagina nao trouxe nada novo (evita laco se o offset for ignorado)
                if not novos:
                    break
                offset += len(dados)
        log(f"{ini} a {fin}: {len(encontrados)} LPCOs acumuladas.")
    return list(encontrados.values())


def carregar_cache():
    if os.path.exists(ARQ_CACHE):
        with open(ARQ_CACHE, encoding="utf-8") as f:
            return json.load(f)
    return {}


def salvar_cache(cache):
    with open(ARQ_CACHE, "w", encoding="utf-8") as f:
        json.dump(cache, f, ensure_ascii=False)


def detalhar_lpcos(api, lpcos, cache):
    """Busca o detalhe de cada LPCO. LPCOs ja em cache e em situacao final nao sao
    consultadas de novo (economiza o limite de acessos por hora)."""
    situacoes_finais = {"DEFERIDO", "INDEFERIDO", "CANCELADO", "VENCIDO", "ANULADO_REVOGADO"}
    for i, l in enumerate(lpcos, 1):
        num = l["numero"]
        em_cache = cache.get(num)
        situacao = re.sub(r"[\s/]+", "_", str(l.get("situacao", "")).strip().upper())
        if em_cache and situacao in situacoes_finais:
            continue
        _, det = api.get(f"/ext/lpco/{num}")
        if det:
            cache[num] = det
        if i % 20 == 0:
            salvar_cache(cache)
            log(f"Detalhadas {i}/{len(lpcos)}")
    salvar_cache(cache)


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def texto_situacao(s):
    if isinstance(s, dict):
        return s.get("descricao") or s.get("codigo") or ""
    return s or ""


def montar_linhas(lpcos, cache, ncms=()):
    """Uma linha por item. Com filtro de NCM, so entram os itens dessas NCMs
    (uma LPCO pode ter itens de outras NCMs tambem)."""
    linhas = []
    for l in sorted(lpcos, key=lambda x: x["numero"]):
        det = cache.get(l["numero"], {})
        base = [
            l["numero"],
            # A consulta traz a situacao de hoje; o detalhe pode estar em cache de outro dia
            l.get("situacao") or texto_situacao(det.get("situacao")),
            l.get("codigoModelo", ""),
            l.get("nome", ""),
            det.get("orgao", ""),
            l.get("dataRegistro", ""),
            l.get("dataInicioVigencia", ""),
            l.get("dataFimVigencia", ""),
            l.get("dueVinculada", ""),
        ]
        itens = det.get("listaNcm") or []
        if ncms:
            itens = [it for it in itens if it.get("ncm") in ncms]
        if not itens:
            linhas.append(base + ["", "", "", "", "", "sem detalhe" if not det else "sem itens"])
            continue
        for it in itens:
            prod = it.get("produto") or {}
            linhas.append(base + [
                it.get("numeroItem", ""),
                it.get("ncm", ""),
                prod.get("codigo", ""),
                prod.get("versao", ""),
                prod.get("cnpjRaiz", ""),
                "" if prod else "item sem produto do catalogo",
            ])
    return linhas


CABECALHO = [
    "Nº LPCO", "Situação", "Modelo", "Nome do modelo", "Órgão", "Data registro",
    "Início vigência", "Fim vigência", "DUIMP/DU-E vinculada",
    "Item", "NCM", "Código produto", "Versão produto", "CNPJ raiz produto", "Observação",
]


def gerar_excel(linhas, caminho, cfg):
    wb = Workbook()
    ws = wb.active
    ws.title = "LPCOs"
    ws.append(CABECALHO)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2E7D32")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for linha in linhas:
        ws.append(linha)
    for col in (11,):  # NCM como texto, sem perder zero a esquerda
        for cel in ws.iter_cols(min_col=col, max_col=col, min_row=2):
            for c in cel:
                c.number_format = "@"
    for i, titulo in enumerate(CABECALHO, 1):
        maior = max([len(str(titulo))] + [len(str(l[i - 1])) for l in linhas[:500]])
        ws.column_dimensions[get_column_letter(i)].width = min(maior + 2, 45)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    info = wb.create_sheet("Controle")
    info.append(["Gerado em", datetime.now().strftime("%d/%m/%Y %H:%M")])
    info.append(["CNPJ", cfg["cnpj"]])
    info.append(["Período", f"{cfg['data_inicial']} a {cfg.get('data_final') or date.today().isoformat()}"])
    info.append(["Ambiente", cfg.get("ambiente", "producao")])
    info.append(["NCMs filtradas", ", ".join(ncms_do_config(cfg)) or "todas"])
    info.append(["LPCOs", len({l[0] for l in linhas})])
    info.append(["Linhas (itens)", len(linhas)])
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 30
    wb.save(caminho)


def main():
    cfg = carregar_config()
    saida = os.path.join(PASTA, cfg.get("arquivo_saida", "LPCOs_Portal_Unico.xlsx"))
    api = PortalUnico(cfg)
    cache = carregar_cache()
    lpcos = []
    try:
        lpcos = listar_lpcos(api, cfg)
        log(f"Total: {len(lpcos)} LPCOs. Buscando detalhes...")
        detalhar_lpcos(api, lpcos, cache)
    except LimiteAtingido as e:
        salvar_cache(cache)
        log(f"LIMITE DO PORTAL ATINGIDO: {e}")
        log("O Excel sera gerado com o que ja foi baixado. Rode de novo depois do horario "
            "indicado: o que ja esta em cache nao e consultado outra vez.")
    except Exception as e:
        salvar_cache(cache)
        log(f"ERRO, extracao interrompida: {type(e).__name__}: {str(e)[:500]}")
        log("O Excel sera gerado com o que ja foi baixado. Rode de novo: "
            "o que ja esta em cache nao e consultado outra vez.")
    finally:
        api.fechar()

    if not lpcos:
        sys.exit("A lista de LPCOs nao foi obtida; o Excel anterior foi mantido.")
    linhas = montar_linhas(lpcos, cache, ncms_do_config(cfg))
    try:
        gerar_excel(linhas, saida, cfg)
    except PermissionError:
        sys.exit(f"Nao consegui salvar {saida}. Feche o arquivo no Excel e rode de novo.")
    log(f"Excel salvo: {saida} ({len(linhas)} linhas)")


if __name__ == "__main__":
    main()
