"""
Segunda extracao: dados do Catalogo de Produtos para os produtos encontrados nas LPCOs.

Le o Excel gerado pelo extrair_lpcos.py (codigos e versoes de produto das LPCOs),
baixa o catalogo da empresa pela API do Portal Unico (catp) e gera um Excel com:
codigo, descricao, descricao detalhada, codigo interno, operador estrangeiro vinculado
e as caracteristicas do pneu lidas da descricao (medida, carga, lonas, marca e modelo).

Usa o mesmo config_lpco.json e a mesma chave de acesso da primeira extracao.
Sao so 3 chamadas por CNPJ raiz (exportacao do catalogo inteiro), entao nao
pesa no limite de acessos do portal.
"""

import io
import json
import os
import sys
import zipfile
from collections import defaultdict
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from extrair_lpcos import PASTA, LimiteAtingido, PortalUnico, carregar_config, log
from leitura_pneus import ler_produto_catalogo


def produtos_das_lpcos(caminho_excel):
    """{(cnpjRaiz, codigo): {"versoes": {...}, "lpcos": {...}, "ncms": {...}}} lido do Excel das LPCOs."""
    if not os.path.exists(caminho_excel):
        sys.exit(f"Nao achei {caminho_excel}. Rode primeiro o extrator de LPCOs.")
    ws = load_workbook(caminho_excel, read_only=True)["LPCOs"]
    linhas = ws.iter_rows(values_only=True)
    cab = {nome: i for i, nome in enumerate(next(linhas))}
    produtos = defaultdict(lambda: {"versoes": set(), "lpcos": set(), "ncms": set()})
    for l in linhas:
        codigo = l[cab["Código produto"]]
        if codigo in (None, ""):
            continue
        raiz = str(l[cab["CNPJ raiz produto"]] or "").zfill(8)
        p = produtos[(raiz, int(codigo))]
        p["versoes"].add(str(l[cab["Versão produto"]] or ""))
        p["lpcos"].add(l[cab["Nº LPCO"]])
        p["ncms"].add(str(l[cab["NCM"]] or ""))
    return produtos


def como_lista(dados):
    if dados is None:
        return []
    return dados if isinstance(dados, list) else [dados]


def exportar(api, caminho, nome):
    """As exportacoes do catalogo vem num .zip com JSON dentro. Guarda o .zip original
    na pasta exportacoes_catalogo/ para conferencia."""
    _, conteudo = api.get(caminho, sistema="catp", binario=True)
    if not conteudo:
        return []
    pasta = os.path.join(PASTA, "exportacoes_catalogo")
    os.makedirs(pasta, exist_ok=True)
    with open(os.path.join(pasta, f"{nome}.zip"), "wb") as f:
        f.write(conteudo)
    if not zipfile.is_zipfile(io.BytesIO(conteudo)):
        return como_lista(json.loads(conteudo.decode("utf-8-sig")))
    registros = []
    with zipfile.ZipFile(io.BytesIO(conteudo)) as z:
        for arq in z.namelist():
            if arq.lower().endswith(".json"):
                registros += como_lista(json.loads(z.read(arq).decode("utf-8-sig")))
        if not registros:
            sys.exit(f"O zip de {nome} nao tem JSON dentro (arquivos: {z.namelist()}). "
                     f"Ele foi salvo em {pasta}; me mande a lista de arquivos.")
    return registros


def baixar_catalogo(api, raiz):
    """Exporta produtos, vinculos de fabricante e operadores estrangeiros de um CNPJ raiz."""
    log(f"Baixando catalogo de produtos do CNPJ raiz {raiz}...")
    prods = exportar(api, f"/ext/produto/exportar/{raiz}/true", f"produtos_{raiz}")
    log("Baixando vinculos produto x fabricante/operador estrangeiro...")
    vinculos = exportar(api, f"/ext/fabricante/exportar/{raiz}", f"fabricantes_{raiz}")
    log("Baixando cadastro de operadores estrangeiros...")
    operadores = exportar(api, f"/ext/operador-estrangeiro/exportar/{raiz}/true", f"operadores_{raiz}")
    log(f"{len(prods)} produtos, {len(vinculos)} vinculos, {len(operadores)} operadores.")
    return prods, vinculos, operadores


def montar_linhas(produtos_lpco, catalogos):
    linhas = []
    for (raiz, codigo), info in sorted(produtos_lpco.items()):
        prods, vinculos, operadores = catalogos.get(raiz, ([], [], []))
        prod = next((p for p in prods if p.get("codigo") == codigo), None)
        ops_por_chave = {(o.get("codigoPais"), str(o.get("codigo"))): o for o in operadores}
        ops = []
        for v in vinculos:
            if v.get("codigoProduto") != codigo:
                continue
            if v.get("codigoOperadorEstrangeiro"):
                o = ops_por_chave.get((v.get("codigoPais"), str(v["codigoOperadorEstrangeiro"])), {})
                ops.append((v["codigoOperadorEstrangeiro"], o.get("nome", ""), v.get("codigoPais", ""), o.get("tin", "")))
            elif v.get("cpfCnpjFabricante"):
                ops.append((v["cpfCnpjFabricante"], "fabricante nacional", "BR", ""))
            else:
                ops.append(("", "fabricante desconhecido", v.get("codigoPais", ""), ""))

        versoes_lpco = ", ".join(sorted(v for v in info["versoes"] if v))
        obs = []
        if not prod:
            obs.append("produto nao encontrado no catalogo")
        elif prod.get("versao") and prod.get("versao") not in info["versoes"]:
            obs.append("versao atual do catalogo diferente da usada na LPCO")
        if prod and not ops:
            obs.append("sem operador estrangeiro vinculado")
        prod = prod or {}
        pneu = ler_produto_catalogo(prod.get("denominacao", ""), prod.get("descricao", ""))
        linhas.append([
            codigo,
            versoes_lpco,
            prod.get("versao", ""),
            prod.get("situacao", ""),
            prod.get("ncm") or ", ".join(sorted(info["ncms"])),
            prod.get("denominacao", ""),
            prod.get("descricao", ""),
            "; ".join(str(c) for c in prod.get("codigosInterno") or []),
            "\n".join(str(o[0]) for o in ops),
            "\n".join(o[1] for o in ops),
            "\n".join(o[2] for o in ops),
            "\n".join(o[3] for o in ops),
            len(info["lpcos"]),
            ", ".join(sorted(info["lpcos"])),
            raiz,
            pneu["medida"],
            pneu["carga"],
            pneu["lonas"],
            pneu["marca"],
            pneu["modelo"],
            "; ".join(obs),
        ])
    return linhas


CABECALHO = [
    "Código produto", "Versão na LPCO", "Versão atual catálogo", "Situação produto", "NCM",
    "Descrição", "Descrição detalhada", "Código interno",
    "Cód. operador estrangeiro", "Nome operador estrangeiro", "País operador", "TIN operador",
    "Qtd LPCOs", "LPCOs", "CNPJ raiz",
    "Medida", "Índice de carga", "Lonas", "Marca", "Modelo", "Observação",
]
LARGURAS = {6: 45, 7: 80, 10: 40, 14: 40}


def gerar_excel(linhas, caminho, origem):
    wb = Workbook()
    ws = wb.active
    ws.title = "Produtos"
    ws.append(CABECALHO)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2E7D32")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for linha in linhas:
        ws.append(linha)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
        row[4].number_format = "@"
    for i, titulo in enumerate(CABECALHO, 1):
        ws.column_dimensions[get_column_letter(i)].width = LARGURAS.get(i, max(len(titulo) + 2, 14))
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    info = wb.create_sheet("Controle")
    info.append(["Gerado em", datetime.now().strftime("%d/%m/%Y %H:%M")])
    info.append(["Origem dos códigos", origem])
    info.append(["Produtos", len(linhas)])
    info.append(["Não encontrados no catálogo", sum(1 for l in linhas if "nao encontrado" in l[-1])])
    info.column_dimensions["A"].width = 28
    info.column_dimensions["B"].width = 60
    wb.save(caminho)


def main():
    cfg = carregar_config()
    entrada = os.path.join(PASTA, cfg.get("arquivo_saida", "LPCOs_Portal_Unico.xlsx"))
    saida = os.path.join(PASTA, cfg.get("arquivo_saida_catalogo", "Catalogo_Produtos_LPCOs.xlsx"))
    produtos_lpco = produtos_das_lpcos(entrada)
    if not produtos_lpco:
        sys.exit("O Excel das LPCOs nao tem nenhum codigo de produto.")
    log(f"{len(produtos_lpco)} produtos distintos encontrados nas LPCOs.")

    api = PortalUnico(cfg)
    catalogos = {}
    try:
        for raiz in sorted({r for r, _ in produtos_lpco}):
            catalogos[raiz] = baixar_catalogo(api, raiz)
    except LimiteAtingido as e:
        sys.exit(f"LIMITE DO PORTAL ATINGIDO: {e}")
    finally:
        api.fechar()

    linhas = montar_linhas(produtos_lpco, catalogos)
    try:
        gerar_excel(linhas, saida, os.path.basename(entrada))
    except PermissionError:
        sys.exit(f"Nao consegui salvar {saida}. Feche o arquivo no Excel e rode de novo.")
    log(f"Excel salvo: {saida} ({len(linhas)} produtos)")


if __name__ == "__main__":
    main()
